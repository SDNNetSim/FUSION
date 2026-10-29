"""TopologyBench parsing and simulation integration without network access."""

import io
import json
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from fusion.core.request import generate_simulation_requests
from fusion.domain.config import SimulationConfig
from fusion.io.topologybench import download_topologybench, list_topologies, load_topologybench
from fusion.sim.input_setup import create_input


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "TOP_1_EXAMPLE.xlsx"
    book = Workbook()
    nodes = book.active
    nodes.title = "Nodes_EXAMPLE"
    nodes.append(["Node_ID", "Latitude", "Longitude", "Location Name", "Country"])
    nodes.append(["1", 33.75, -84.3833, "Atlanta", "US"])
    nodes.append(["2", 40, -80, "Example", "US"])
    nodes.append(["03", None, None, None, None])
    edges = book.create_sheet("Edges_EXAMPLE")
    edges.append(["Edge_ID", "Source", "Destination", "Computed Length (km)", "Lookup"])
    edges.append([1, 1, 2, 885.1138413916553])
    edges.append([2, 2, "03", 15.25])
    edges.append([None, None, None, None, "auxiliary lookup, not an edge"])
    book.create_sheet("Location Info")
    book.save(path)
    book.close()
    return path


def edit_cell(path, sheet, cell, value):
    book = load_workbook(path)
    book[sheet][cell] = value
    book.save(path)
    book.close()


def test_preserves_ids_exact_distances_and_node_metadata(workbook):
    result = load_topologybench(workbook)
    assert list(result.nodes) == ["1", "2", "03"]
    assert result.links == {("1", "2"): 885.1138413916553, ("2", "03"): 15.25}
    assert result.nodes["1"]["location_name"] == "Atlanta"
    assert result.nodes["1"]["latitude"] == 33.75
    assert result.summary["num_links"] == 2
    assert len(result.summary["workbook_sha256"]) == 64


def test_constant_weight_is_explicit_and_records_original_length(workbook):
    result = load_topologybench(workbook, constant_weight=True)
    assert set(result.links.values()) == {1.0}
    assert result.summary["source_total_length_km"] == 900.3638413916553
    assert result.summary["constant_weight"] is True


def test_archive_selection_is_unambiguous_and_does_not_extract(workbook, tmp_path):
    path = tmp_path / "dataset.zip"
    with ZipFile(path, "w") as archive:
        archive.write(workbook, "real/TOP_1_EXAMPLE.xlsx")
        archive.write(workbook, "real/TOP_2_EXAMPLE.xlsx")
        archive.writestr("../../escape.txt", "not extracted")
    assert len(list_topologies(path)) == 2
    for name in ("", "EXAMPLE", "missing"):
        with pytest.raises(ValueError, match="Select exactly one"):
            load_topologybench(path, name)
    for name in ("TOP_1_EXAMPLE", "TOP_1_EXAMPLE.xlsx", "real/TOP_1_EXAMPLE.xlsx"):
        assert load_topologybench(path, name).links == load_topologybench(workbook).links
    assert not (tmp_path / "real").exists()


def test_single_member_archive_does_not_require_name(workbook, tmp_path):
    path = tmp_path / "single.zip"
    with ZipFile(path, "w") as archive:
        archive.write(workbook, workbook.name)
    assert load_topologybench(path).summary["num_nodes"] == 3


@pytest.mark.parametrize(
    "sheet,cell,value,message",
    [
        ("Nodes_EXAMPLE", "A3", "1", "Duplicate node"),
        ("Nodes_EXAMPLE", "A3", None, "Missing or invalid"),
        ("Nodes_EXAMPLE", "B3", 91, "Invalid latitude"),
        ("Edges_EXAMPLE", "C2", "missing", "Unknown edge endpoint"),
        ("Edges_EXAMPLE", "C2", 1, "Self-link"),
        ("Edges_EXAMPLE", "D2", 0, "positive and finite"),
        ("Edges_EXAMPLE", "D2", -1, "positive and finite"),
        ("Edges_EXAMPLE", "D2", "nan", "positive and finite"),
        ("Edges_EXAMPLE", "D2", "inf", "positive and finite"),
        ("Edges_EXAMPLE", "D2", "=100+5", "Missing or invalid computed length"),
        ("Edges_EXAMPLE", "D1", "Length", "Unexpected edge columns"),
    ],
)
def test_invalid_data_fails_before_simulation(workbook, sheet, cell, value, message):
    edit_cell(workbook, sheet, cell, value)
    with pytest.raises(ValueError, match=message):
        load_topologybench(workbook)


def test_reverse_duplicate_is_not_silently_collapsed(workbook):
    edit_cell(workbook, "Edges_EXAMPLE", "C3", 1)
    with pytest.raises(ValueError, match="Duplicate/parallel"):
        load_topologybench(workbook)


def test_isolated_node_is_not_silently_dropped(workbook):
    book = load_workbook(workbook)
    book["Nodes_EXAMPLE"].append(["4"])
    book.save(workbook)
    book.close()
    with pytest.raises(ValueError, match="connected topology"):
        load_topologybench(workbook)


def test_rejects_pickle_without_opening_it(tmp_path):
    with pytest.raises(ValueError, match="pickle is not supported"):
        load_topologybench(tmp_path / "graphs.pkl")


def test_workbook_name_mismatch(workbook):
    with pytest.raises(ValueError, match="does not match"):
        load_topologybench(workbook, "different")


def test_download_verifies_content_before_saving(tmp_path):
    path = tmp_path / "download.zip"
    with patch("fusion.io.topologybench.urllib.request.urlopen", return_value=io.BytesIO(b"invalid download")):
        with pytest.raises(ValueError, match="checksum"):
            download_topologybench(path)
    assert not path.exists()


def test_download_never_overwrites_existing_file(tmp_path):
    path = tmp_path / "download.zip"
    path.write_bytes(b"existing")
    with patch("fusion.io.topologybench.urllib.request.urlopen") as fetch:
        with pytest.raises(FileExistsError):
            download_topologybench(path)
        fetch.assert_not_called()
    assert path.read_bytes() == b"existing"


@pytest.mark.parametrize("mode", ["dynamic", "static", "semi_static"])
def test_import_to_physical_topology_and_request_generator(workbook, tmp_path, mode):
    from fusion.cli.config_setup import load_config

    template = Path(__file__).resolve().parents[2] / "configs/templates/static_design.ini"
    props = load_config(str(template))["s1"]
    props.update(
        topology_file=str(workbook),
        network="Example",
        num_requests=10,
        traffic_type=mode,
        request_sort="none" if mode == "dynamic" else "bandwidth",
        arrival_rate=1,
        thread_num="s1",
        date="test",
        sim_start=mode,
        is_only_core_node=False,
    )
    prepared = create_input(str(tmp_path), props)
    assert set(prepared["topology_info"]["nodes"]) == {"1", "2", "03"}
    assert prepared["topology_info"]["nodes"]["1"]["location_name"] == "Atlanta"
    assert prepared["topology_info"]["links"][1]["length"] == 885.1138413916553
    assert prepared["topology_import_summary"]["num_links"] == 2
    assert set(prepared["core_nodes"]) == {"1", "2", "03"}
    saved = json.loads((tmp_path / "input" / "Example" / "test" / mode / "topology_import_s1.json").read_text())
    assert saved["summary"] == prepared["topology_import_summary"]
    assert saved["topology_info"]["links"]["1"]["length"] == 885.1138413916553
    requests = generate_simulation_requests(seed=42, engine_properties=prepared)
    assert len(requests) == (20 if mode == "dynamic" else 10)
    restored = SimulationConfig.from_engine_props(prepared).to_engine_props()
    assert restored["topology_file"] == str(workbook)
    assert restored["topology_info"] == prepared["topology_info"]
