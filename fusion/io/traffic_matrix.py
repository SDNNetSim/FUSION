"""Load finite demand sets from CSV, JSON, square CSV matrices, or SNDlib XML.

Amounts use the same units as mod_per_bw (normally Gbit/s). Download public
repository files first so experiments use an immutable local input.
"""

import argparse
import csv
import io
import json
import math
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

from fusion.utils.bandwidth import bandwidth_key

FORMATS = ("auto", "csv", "json", "matrix_csv", "matrix_xlsx", "sndlib")
MAX_DOWNLOAD_BYTES = 20 * 1024 * 1024


def _positive_number(value: Any, name: str, *, allow_zero: bool = False) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{name} must be a finite number: {value!r}") from exc
    if not math.isfinite(number) or number < 0 or (number == 0 and not allow_zero):
        raise ValueError(f"{name} must be finite and {'non-negative' if allow_zero else 'positive'}")
    return number


def _matrix_rows(text: str, direction: str = "directed") -> list[dict[str, Any]]:
    table = list(csv.reader(io.StringIO(text)))
    if len(table) < 2:
        raise ValueError("Matrix CSV needs a header and at least one data row")
    nodes = [cell.strip() for cell in table[0][1:]]
    if not nodes or any(not node for node in nodes) or len(set(nodes)) != len(nodes):
        raise ValueError("Matrix CSV column node labels must be nonempty and unique")
    if len(table) != len(nodes) + 1 or any(len(row) != len(nodes) + 1 for row in table[1:]):
        raise ValueError("Matrix CSV must be square with row and column labels")
    sources = [row[0].strip() for row in table[1:]]
    if len(set(sources)) != len(nodes) or set(sources) != set(nodes):
        raise ValueError("Matrix CSV row labels must match column labels exactly")
    rows = []
    cells = {}
    for source, row in zip(sources, table[1:], strict=True):
        for destination, value in zip(nodes, row[1:], strict=True):
            amount = _positive_number(value, "Matrix cell", allow_zero=True)
            if source == destination and amount:
                raise ValueError("Matrix CSV diagonal must be zero")
            if amount:
                cells[(source, destination)] = amount
    if direction not in ("directed", "undirected"):
        raise ValueError("traffic_matrix_direction must be directed or undirected")
    for (source, destination), amount in cells.items():
        if direction == "undirected":
            if not math.isclose(amount, cells.get((destination, source), 0), rel_tol=1e-9, abs_tol=1e-9):
                raise ValueError("Undirected traffic requires a symmetric matrix")
            if nodes.index(source) > nodes.index(destination):
                continue
        rows.append({"source": source, "destination": destination, "bandwidth": amount})
    return rows


def load_excel_matrix(path: str, sheet: str, direction: str = "directed") -> list[dict[str, Any]]:
    """Read a labelled matrix, including TEFNET24's padded worksheets and blank zeros."""
    from openpyxl import load_workbook

    with Path(path).open("rb") as source:
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            if not sheet or sheet not in workbook.sheetnames:
                raise ValueError(f"Set traffic_matrix_sheet to one of {workbook.sheetnames}")
            raw_rows = list(workbook[sheet].iter_rows(values_only=True))
        finally:
            workbook.close()
    while raw_rows and all(value in (None, "") for value in raw_rows[-1]):
        raw_rows.pop()
    if not raw_rows:
        raise ValueError("Excel matrix sheet is empty")
    header = list(raw_rows[0])
    while header and header[-1] in (None, ""):
        header.pop()
    table = [header]
    for row in raw_rows[1:]:
        if any(value not in (None, "") for value in row[len(header) :]):
            raise ValueError("Excel matrix has data beyond its column labels")
        table.append([row[0], *(0 if value in (None, "") else value for value in row[1 : len(header)])])
    output = io.StringIO()
    csv.writer(output).writerows(table)
    return _matrix_rows(output.getvalue(), direction)


def _sndlib_rows(text: str) -> list[dict[str, Any]]:
    # Reject DTDs/entities; these are unnecessary for SNDlib network instances.
    if "<!DOCTYPE" in text.upper() or "<!ENTITY" in text.upper():
        raise ValueError("SNDlib XML must not contain DTDs or entity declarations")
    try:
        root = ET.fromstring(text)
    except ET.ParseError as exc:
        raise ValueError(f"Invalid SNDlib XML: {exc}") from exc
    rows = []
    for demand in root.findall(".//{*}demands/{*}demand"):
        fields = {child.tag.rsplit("}", 1)[-1]: child.text for child in demand}
        amount = _positive_number(fields.get("demandValue"), "SNDlib demandValue")
        unit = _positive_number(fields.get("routingUnit", 1), "SNDlib routingUnit")
        rows.append({"source": fields.get("source"), "destination": fields.get("target"), "bandwidth": amount * unit})
    return rows


def load_traffic_matrix(
    path: str, file_format: str = "auto", scale: float = 1.0, sheet: str = "", direction: str = "directed"
) -> list[dict[str, Any]]:
    """Read demands, expanding optional integer counts and validating timestamps.

    CSV/JSON records require source, destination, bandwidth; optional count and
    arrival fields specify multiplicity and explicit arrival times. All records
    must either supply arrival or omit it. Duplicate pairs are independent demands.
    """
    scale = _positive_number(scale, "traffic_matrix_scale")
    if file_format not in FORMATS:
        raise ValueError(f"traffic_matrix_format must be one of {FORMATS}")
    file_path = Path(path).expanduser()
    if file_format == "auto":
        file_format = {".csv": "csv", ".json": "json", ".xml": "sndlib", ".xlsx": "matrix_xlsx"}.get(file_path.suffix.lower(), "")
        if not file_format:
            raise ValueError("Cannot infer traffic matrix format; set traffic_matrix_format explicitly")
    text = "" if file_format == "matrix_xlsx" else file_path.read_text(encoding="utf-8-sig")
    if file_format == "matrix_xlsx":
        rows = load_excel_matrix(str(file_path), sheet, direction)
    elif file_format == "json":
        rows = json.loads(text)
        if isinstance(rows, dict):
            rows = rows.get("demands")
    elif file_format == "csv":
        reader = csv.DictReader(io.StringIO(text))
        if not {"source", "destination", "bandwidth"}.issubset(reader.fieldnames or []):
            raise ValueError("Demand CSV needs source,destination,bandwidth headers; use matrix_csv for square matrices")
        rows = list(reader)
    elif file_format == "matrix_csv":
        rows = _matrix_rows(text, direction)
    else:
        rows = _sndlib_rows(text)
    if not isinstance(rows, list) or not rows:
        raise ValueError("Traffic matrix must contain at least one demand")
    demands = []
    for index, row in enumerate(rows, 1):
        if not isinstance(row, dict):
            raise ValueError(f"Demand {index} must be an object")
        endpoints = []
        for field in ("source", "destination"):
            value = row.get(field)
            if value is None or not str(value).strip():
                raise ValueError(f"Demand {index}: missing {field}")
            endpoints.append(str(value).strip())
        if endpoints[0] == endpoints[1]:
            raise ValueError(f"Demand {index}: source and destination must differ")
        amount = _positive_number(row.get("bandwidth"), f"Demand {index} bandwidth") * scale
        amount = _positive_number(amount, f"Demand {index} scaled bandwidth")
        count_value = row.get("count", 1)
        count = _positive_number(1 if count_value in (None, "") else count_value, f"Demand {index} count")
        if not count.is_integer():
            raise ValueError(f"Demand {index}: count must be a positive integer")
        demand = {"source": endpoints[0], "destination": endpoints[1], "bandwidth": bandwidth_key(amount)}
        if row.get("arrival") not in (None, ""):
            demand["arrival"] = _positive_number(row["arrival"], f"Demand {index} arrival", allow_zero=True)
        demands.extend(dict(demand) for _ in range(int(count)))
    if any("arrival" in row for row in demands) and not all("arrival" in row for row in demands):
        raise ValueError("Either all demands must specify arrival times, or none may")
    return demands


def prepare_traffic_matrix(engine_props: dict[str, Any]) -> None:
    """Load once per run, validating topology/bandwidth and updating statistics inputs."""
    path = engine_props.get("traffic_matrix_file")
    if not path:
        return
    if "traffic_demands" not in engine_props:
        demands = load_traffic_matrix(
            path,
            engine_props.get("traffic_matrix_format", "auto"),
            engine_props.get("traffic_matrix_scale", 1.0),
            engine_props.get("traffic_matrix_sheet", ""),
            engine_props.get("traffic_matrix_direction", "directed"),
        )
        offered_bandwidth = sum(float(row["bandwidth"]) for row in demands)
        node_map_path = engine_props.get("traffic_node_map_file")
        node_map = json.loads(Path(node_map_path).read_text(encoding="utf-8")) if node_map_path else {}
        if not isinstance(node_map, dict):
            raise ValueError("traffic_node_map_file must contain a JSON object mapping input node names to topology node names")
        nodes = {str(node): node for node in engine_props["topology_info"]["nodes"]}
        bandwidths = {float(bw): bw for bw in engine_props["mod_per_bw"]}
        for row in demands:
            for field in ("source", "destination"):
                mapped = str(node_map.get(row[field], row[field]))
                if mapped not in nodes:
                    raise ValueError(f"Traffic node {mapped!r} is absent from the topology; set traffic_node_map_file if needed")
                row[field] = nodes[mapped]
            if row["source"] == row["destination"]:
                raise ValueError("Node mapping turns a demand into a self-demand")
            amount = float(row["bandwidth"])
            if amount not in bandwidths:
                raise ValueError(
                    f"Traffic bandwidth {amount:g} is not supported by mod_per_bw; "
                    "adjust traffic_matrix_scale, preprocess demands, or add modulation definitions"
                )
            row["bandwidth"] = bandwidths[amount]
        engine_props["traffic_demands"] = demands
        engine_props["traffic_matrix_summary"] = {
            "offered_bandwidth": offered_bandwidth,
            "simulated_bandwidth": sum(float(row["bandwidth"]) for row in demands),
            "num_requests": len(demands),
        }
    demands = engine_props["traffic_demands"]
    engine_props["num_requests"] = len(demands)
    counts = dict.fromkeys(engine_props["mod_per_bw"], 0)
    for row in demands:
        counts[row["bandwidth"]] += 1
    engine_props["request_distribution"] = {bw: count / len(demands) for bw, count in counts.items()}


def download_traffic_matrix(url: str, destination: str) -> Path:
    """Save a public HTTPS raw file without overwriting an existing experiment input."""
    if urllib.parse.urlparse(url).scheme != "https":
        raise ValueError("Use an HTTPS raw-file URL")
    with urllib.request.urlopen(url, timeout=30) as response:  # nosec B310: HTTPS checked above
        if urllib.parse.urlparse(response.geturl()).scheme != "https":
            raise ValueError("Traffic download redirected away from HTTPS")
        data = response.read(MAX_DOWNLOAD_BYTES + 1)
    if len(data) > MAX_DOWNLOAD_BYTES:
        raise ValueError("Traffic matrix download exceeds 20 MiB")
    destination_path = Path(destination)
    with destination_path.open("xb") as output:
        output.write(data)
    return destination_path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Download a public traffic file for reproducible local simulation")
    parser.add_argument("url", help="HTTPS raw CSV/JSON or SNDlib XML URL (prefer a pinned repository commit)")
    parser.add_argument("destination", help="New local file path")
    args = parser.parse_args()
    print(download_traffic_matrix(args.url, args.destination))
